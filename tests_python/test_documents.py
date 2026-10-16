from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from featurelens.core import Service
from featurelens.tools import invoke


@pytest.fixture
def setup(tmp_path):
    svc = Service(tmp_path / "data")
    s = svc.create("文档", "证据？")
    return svc, s, tmp_path


def test_text_locations_capture_and_integrity(setup):
    svc, s, p = setup
    path = p / "notes.md"
    path.write_text("# 规则\n\n无资料时提示不足\n")
    d = invoke(svc, "research_import_document", dict(study_id=s["id"], path=str(path)))
    assert invoke(
        svc, "research_import_document", dict(study_id=s["id"], path=str(path))
    )["reused"]
    result = invoke(
        svc,
        "research_document_chunks",
        dict(study_id=s["id"], document_id=d["id"], query="不足"),
    )
    chunk = result["items"][0]
    assert chunk["location"]["line"] == 3
    e = invoke(
        svc,
        "research_capture_document_chunk",
        dict(
            study_id=s["id"],
            document_id=d["id"],
            chunk_id=chunk["chunk_id"],
            title="无资料规则",
        ),
    )
    assert e["source_sha256"] == d["sha256"]
    assert e["location"]["line"] == 3
    svc.claim(s["id"], "文档包含无资料规则", "fact", [e["id"]])
    assert "无资料时提示不足" in svc.report(s["id"])["text"]
    Path(d["chunks_file"]).write_text("[]")
    with pytest.raises(ValueError, match="片段已变化"):
        invoke(
            svc, "research_document_chunks", dict(study_id=s["id"], document_id=d["id"])
        )


def test_docx_body_and_table(setup):
    svc, s, p = setup
    doc = Document()
    doc.add_paragraph("产品规则")
    doc.add_paragraph("权限不足时停止")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "条件"
    table.cell(0, 1).text = "结果"
    path = p / "test.docx"
    doc.save(path)
    d = invoke(svc, "research_import_document", dict(study_id=s["id"], path=str(path)))
    result = invoke(
        svc, "research_document_chunks", dict(study_id=s["id"], document_id=d["id"])
    )
    assert result["items"][1]["location"]["paragraph"] == 2
    assert result["items"][2]["location"]["table"] == 1
    assert d["warnings"]


def test_pdf_page_locations_and_blank_pages(setup):
    svc, s, p = setup
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (No evidence found.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(300, 300)
    path = p / "test.pdf"
    writer.write(path)
    d = invoke(svc, "research_import_document", dict(study_id=s["id"], path=str(path)))
    result = invoke(
        svc, "research_document_chunks", dict(study_id=s["id"], document_id=d["id"])
    )
    assert result["items"][0]["location"]["page"] == 1
    assert "No evidence" in result["items"][0]["text"]
    assert any("第 2 页" in w for w in d["warnings"])
    blank = PdfWriter()
    blank.add_blank_page(300, 300)
    empty = p / "blank.pdf"
    blank.write(empty)
    with pytest.raises(ValueError, match="未提取到文字"):
        invoke(svc, "research_import_document", dict(study_id=s["id"], path=str(empty)))


def test_unsupported_and_cross_study(setup):
    svc, s, p = setup
    file = p / "app.exe"
    file.write_bytes(b"hello")
    with pytest.raises(ValueError, match="只支持"):
        invoke(svc, "research_import_document", dict(study_id=s["id"], path=str(file)))
    with pytest.raises(ValueError, match="绝对路径"):
        invoke(
            svc, "research_import_document", dict(study_id=s["id"], path="relative.txt")
        )
    file = p / "ok.txt"
    file.write_text("sample")
    d = invoke(svc, "research_import_document", dict(study_id=s["id"], path=str(file)))
    other = svc.create("其他", "问题")
    with pytest.raises(ValueError, match="找不到文档"):
        invoke(
            svc,
            "research_document_chunks",
            dict(study_id=other["id"], document_id=d["id"]),
        )
    Path(d["raw_file"]).write_text("tampered")
    with pytest.raises(ValueError, match="原始文档已变化"):
        invoke(
            svc, "research_document_chunks", dict(study_id=s["id"], document_id=d["id"])
        )
